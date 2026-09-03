"""Full-text search over the documents in the data folder.

Why this exists: without it, a question like "find the invoice for the sensor
calibration" forces the model to open PDFs one at a time looking for the words.
Each one is a round trip, and MAX_TOOL_STEPS caps that at about nine files. Fine
for a test folder, useless at two hundred.

Everything here is a CACHE. Every word in the index was derived from a file you
already have, so deleting the whole thing costs nothing but the time to rebuild.
Four rules keep it that way, and they are what make moving to PostgreSQL later a
change of storage rather than a migration:

  1. search(), index_file() and rebuild() are the only interface. Nothing else
     in the app knows SQLite exists.
  2. The extracted TEXT is stored, not only the inverted index. Migrating is
     then an insert, not a re-parse of every PDF.
  3. The database lives outside DATA_DIR, so it is obviously not user data.
  4. rebuild() reconstructs it from scratch, which makes it disposable by
     construction rather than by intention.

Nothing durable is ever stored here. No tags, no notes, no annotations. The
moment something exists only in this file it stops being a cache.
"""

from __future__ import annotations

import re
import sqlite3
import time
from pathlib import Path

from app.config import INDEX_PATH, ensure_data_dir, ensure_index_dir

SEARCHABLE = {".pdf", ".xlsx", ".xlsm"}
MAX_TEXT_PER_FILE = 400_000


class SearchError(Exception):
    """Something the model can read and act on."""


# --------------------------------------------------------------------------
# storage
# --------------------------------------------------------------------------

SCHEMA = """
CREATE VIRTUAL TABLE IF NOT EXISTS documents USING fts5(
    name,
    text,
    size UNINDEXED,
    mtime UNINDEXED,
    indexed_at UNINDEXED,
    tokenize = 'porter unicode61'
);
"""


def connect() -> sqlite3.Connection:
    ensure_index_dir()
    connection = sqlite3.connect(INDEX_PATH, timeout=15)
    connection.row_factory = sqlite3.Row
    # WAL lets several worker processes on one machine read while one writes,
    # which is what a multi-card single server needs. It relies on shared memory
    # and so fails on some network and container mounts; falling back keeps the
    # feature working rather than failing at import time.
    try:
        connection.execute("PRAGMA journal_mode=WAL")
        connection.execute("PRAGMA synchronous=NORMAL")
    except sqlite3.OperationalError:
        pass
    try:
        connection.executescript(SCHEMA)
    except sqlite3.OperationalError as exc:
        connection.close()
        raise SearchError(
            f"Could not open the search index at {INDEX_PATH}: {exc}. "
            "Set INDEX_DIR in .env to a local disk; network shares and some "
            "container mounts cannot host a SQLite database."
        ) from exc
    return connection


# --------------------------------------------------------------------------
# extraction
# --------------------------------------------------------------------------

def extract(path: Path) -> str:
    """Pull readable text out of one file. Empty string if there is none."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            import pymupdf

            with pymupdf.open(path) as document:
                pages = [document.load_page(i).get_text("text") for i in range(document.page_count)]
            return "\n".join(pages)[:MAX_TEXT_PER_FILE]

        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            book = load_workbook(path, data_only=True, read_only=True)
            try:
                parts: list[str] = []
                for sheet in book.worksheets:
                    parts.append(sheet.title)
                    for row in sheet.iter_rows(values_only=True):
                        cells = [str(v) for v in row if v is not None]
                        if cells:
                            parts.append(" ".join(cells))
                        if sum(len(p) for p in parts) > MAX_TEXT_PER_FILE:
                            break
            finally:
                book.close()
            return "\n".join(parts)[:MAX_TEXT_PER_FILE]
    except Exception:  # noqa: BLE001 - an unreadable file is skipped, never fatal
        return ""
    return ""


# --------------------------------------------------------------------------
# writing
# --------------------------------------------------------------------------

def index_file(path: Path, connection: sqlite3.Connection | None = None) -> dict:
    """Add or replace one file in the index. Safe to call repeatedly."""
    own = connection is None
    connection = connection or connect()
    try:
        if path.suffix.lower() not in SEARCHABLE or not path.is_file():
            return {"name": path.name, "indexed": False, "reason": "not a searchable file"}
        text = extract(path)
        stat = path.stat()
        connection.execute("DELETE FROM documents WHERE name = ?", (path.name,))
        if text.strip():
            connection.execute(
                "INSERT INTO documents (name, text, size, mtime, indexed_at) VALUES (?,?,?,?,?)",
                (path.name, text, stat.st_size, stat.st_mtime, time.time()),
            )
        connection.commit()
        return {
            "name": path.name,
            "indexed": bool(text.strip()),
            "characters": len(text),
            "reason": "" if text.strip() else "no extractable text, probably a scan",
        }
    finally:
        if own:
            connection.close()


def remove(name: str) -> None:
    connection = connect()
    try:
        connection.execute("DELETE FROM documents WHERE name = ?", (name,))
        connection.commit()
    finally:
        connection.close()


def rebuild(report=None) -> dict:
    """Throw the index away and build it again from the files on disk.

    This is the proof that the index is disposable. If it can always be
    rebuilt, nothing irreplaceable can accumulate in it.
    """
    folder = ensure_data_dir()
    files = sorted(
        p for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() in SEARCHABLE and not p.name.startswith(".")
    )
    connection = connect()
    started = time.time()
    indexed, skipped = [], []
    try:
        connection.execute("DELETE FROM documents")
        connection.commit()
        for number, path in enumerate(files, start=1):
            result = index_file(path, connection)
            (indexed if result["indexed"] else skipped).append(result)
            if report:
                report(number / max(1, len(files)), f"{number} of {len(files)}: {path.name}")
    finally:
        connection.close()
    return {
        "files_seen": len(files),
        "indexed": len(indexed),
        "skipped": [s["name"] for s in skipped],
        "seconds": round(time.time() - started, 2),
    }


def forget_missing(connection: sqlite3.Connection | None = None) -> list[str]:
    """Drop index rows for files that are no longer on disk.

    The index is written when a file arrives and never when one leaves, so a
    deleted document went on being returned by search_files for ever. The model
    then called read_pdf on it, got "no file named ...", and spent the rest of
    its budget working out that the search result had lied to it. Cheap to run
    on the read path, and it keeps the cache honest without a sweeper.
    """
    own = connection is None
    connection = connection or connect()
    try:
        folder = ensure_data_dir()
        gone = [
            row["name"]
            for row in connection.execute("SELECT name FROM documents")
            if not (folder / row["name"]).is_file()
        ]
        for name in gone:
            connection.execute("DELETE FROM documents WHERE name = ?", (name,))
        if gone:
            connection.commit()
        return gone
    except sqlite3.Error:  # a tidy-up is never worth failing a search over
        return []
    finally:
        if own:
            connection.close()


def stale(connection: sqlite3.Connection | None = None) -> list[str]:
    """Files on disk that the index has not seen, or has seen an older copy of."""
    own = connection is None
    connection = connection or connect()
    try:
        known = {
            row["name"]: (row["size"], row["mtime"])
            for row in connection.execute("SELECT name, size, mtime FROM documents")
        }
        out = []
        for path in ensure_data_dir().iterdir():
            if not path.is_file() or path.suffix.lower() not in SEARCHABLE:
                continue
            if path.name.startswith("."):
                continue
            stat = path.stat()
            seen = known.get(path.name)
            if seen is None or int(seen[0]) != stat.st_size or abs(float(seen[1]) - stat.st_mtime) > 1:
                out.append(path.name)
        return sorted(out)
    finally:
        if own:
            connection.close()


# --------------------------------------------------------------------------
# reading
# --------------------------------------------------------------------------

WORD = re.compile(r"[A-Za-z0-9_][A-Za-z0-9_'-]*")


def _terms(query: str) -> list[str]:
    """Turn a plain question into terms FTS5 will accept.

    The model writes this string, so it can contain quotes, brackets and
    operators that are valid English and invalid FTS5. Extracting words and
    quoting them is safer than passing it through and catching the error.
    """
    stop = {"the", "a", "an", "of", "for", "in", "on", "to", "and", "is", "was",
            "what", "which", "who", "find", "me", "my", "any", "all", "with"}
    words = [w for w in WORD.findall(query or "") if len(w) > 1]
    kept = [w for w in words if w.lower() not in stop] or words
    return [f'"{w}"' for w in kept[:12]]


def search(query: str, limit: int = 8) -> dict:
    """Find documents by what is inside them.

    Tries all terms together first, which is precise. If that finds nothing,
    falls back to any term, which is forgiving. Reports which one answered so
    the model knows how much to trust the match.
    """
    terms = _terms(query)
    if not terms:
        raise SearchError("Nothing searchable in that query. Give me some words to look for.")

    connection = connect()
    try:
        forget_missing(connection)
        total = connection.execute("SELECT count(*) AS n FROM documents").fetchone()["n"]
        for joiner, precision in ((" AND ", "all terms"), (" OR ", "any term")):
            expression = joiner.join(terms)
            try:
                rows = connection.execute(
                    """
                    SELECT name,
                           snippet(documents, 1, '[', ']', ' ... ', 18) AS snippet,
                           bm25(documents) AS score
                    FROM documents
                    WHERE documents MATCH ?
                    ORDER BY score
                    LIMIT ?
                    """,
                    (expression, max(1, min(int(limit), 25))),
                ).fetchall()
            except sqlite3.OperationalError as exc:
                raise SearchError(f"Could not run that search: {exc}") from exc
            if rows:
                return {
                    "query": query,
                    "matched_on": precision,
                    "documents_indexed": total,
                    "count": len(rows),
                    # Said out loud because a list of hits reads like a list of
                    # answers, and it is not. These documents contain the words.
                    # Whether they satisfy a condition is a separate question.
                    "what_this_means": (
                        "These documents CONTAIN these words. This is a text match, not a "
                        "filter: it cannot compare numbers, amounts or dates. If the question "
                        "had a condition in it, read each candidate and check the condition "
                        "yourself before listing it as an answer."
                    ),
                    "results": [
                        {
                            "name": row["name"],
                            "read_with": "read_pdf" if row["name"].lower().endswith(".pdf")
                                         else "read_excel",
                            "snippet": " ".join((row["snippet"] or "").split()),
                            # bm25 returns negative numbers, lower being better.
                            "score": round(-row["score"], 2),
                        }
                        for row in rows
                    ],
                }
        return {
            "query": query,
            "documents_indexed": total,
            "count": 0,
            "results": [],
            "note": (
                "Nothing matched. The index holds "
                f"{total} document(s). Either the words are not in any of them, or a "
                "file has not been indexed yet. list_files shows everything present."
            ),
        }
    finally:
        connection.close()


def status() -> dict:
    connection = connect()
    try:
        row = connection.execute(
            "SELECT count(*) AS n, sum(length(text)) AS chars FROM documents"
        ).fetchone()
        return {
            "index_path": str(INDEX_PATH),
            "documents": row["n"] or 0,
            "characters": row["chars"] or 0,
            "not_yet_indexed": stale(connection),
            "disposable": True,
        }
    finally:
        connection.close()
