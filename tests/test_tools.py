"""Unit tests for the tool functions. Run with:  pytest -q

These use a throwaway data folder, so they never touch your real files.
"""

from __future__ import annotations

import pytest

from app import config, tools


@pytest.fixture(autouse=True)
def temp_data_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    yield tmp_path


# --- path safety ----------------------------------------------------------

@pytest.mark.parametrize(
    "bad_name",
    [
        "../escape.pdf",
        "../../etc/passwd",
        r"..\..\Windows\System32\config",
        "C:/Windows/win.ini",
        "sub/../../out.pdf",
    ],
)
def test_paths_outside_data_dir_are_refused(bad_name):
    with pytest.raises(tools.ToolError):
        tools.read_pdf(bad_name)


def test_a_plain_name_resolves_inside_the_data_dir(temp_data_dir):
    assert config.resolve_in_data_dir("report.pdf").parent == temp_data_dir


def test_a_subfolder_is_allowed(temp_data_dir):
    resolved = config.resolve_in_data_dir("2026/q1.xlsx")
    assert resolved.parent.parent == temp_data_dir


# --- excel ----------------------------------------------------------------

def test_write_then_read_excel_round_trip():
    tools.write_excel(
        "book.xlsx",
        rows=[["a", 1], ["b", 2]],
        headers=["label", "value"],
        mode="overwrite",
    )
    result = tools.read_excel("book.xlsx")
    assert result["likely_header"] == ["label", "value"]
    assert result["rows"][1] == ["a", 1]
    assert result["total_rows"] == 3


def test_append_adds_to_the_end():
    tools.write_excel("book.xlsx", rows=[["a", 1]], headers=["label", "value"])
    result = tools.write_excel("book.xlsx", rows=[["b", 2]], mode="append")
    assert result["total_rows_now"] == 3
    assert result["action"] == "appended to"
    assert tools.read_excel("book.xlsx")["rows"][-1] == ["b", 2]


def test_append_to_a_missing_file_is_a_clear_error():
    with pytest.raises(tools.ToolError, match="does not exist"):
        tools.write_excel("nope.xlsx", rows=[["a"]], mode="append")


def test_named_sheet_is_honoured_and_unknown_sheet_is_reported():
    tools.write_excel("book.xlsx", rows=[["x"]], sheet="Data")
    tools.write_excel("book.xlsx", rows=[["y"]], sheet="Other")
    assert tools.read_excel("book.xlsx", sheet="Data")["sheet"] == "Data"
    # with more than one sheet there is no safe guess, so it must report
    with pytest.raises(tools.ToolError, match="No sheet named"):
        tools.read_excel("book.xlsx", sheet="Missing")


def test_bad_mode_and_bad_suffix_are_refused():
    with pytest.raises(tools.ToolError, match="overwrite"):
        tools.write_excel("book.xlsx", rows=[["a"]], mode="sideways")
    with pytest.raises(tools.ToolError, match=".xlsx"):
        tools.write_excel("book.csv", rows=[["a"]])


def test_rows_must_be_a_list_of_lists():
    with pytest.raises(tools.ToolError):
        tools.write_excel("book.xlsx", rows=["not", "rows"])


def test_max_rows_truncates_and_says_so():
    tools.write_excel("big.xlsx", rows=[[i] for i in range(50)])
    result = tools.read_excel("big.xlsx", max_rows=10)
    assert result["rows_returned"] == 10
    assert result["truncated"] is True
    assert result["total_rows"] == 50


# --- pdf ------------------------------------------------------------------

def test_write_then_read_pdf_round_trip():
    tools.write_pdf("doc.pdf", title="Title Here", body="The magic number is 4815162342.")
    result = tools.read_pdf("doc.pdf")
    assert "4815162342" in result["text"]
    assert result["page_count"] == 1


def test_page_selection():
    body = "\n\n".join(f"Paragraph {i}" for i in range(400))
    tools.write_pdf("long.pdf", title="Long", body=body)
    full = tools.read_pdf("long.pdf")
    assert full["page_count"] > 1
    first = tools.read_pdf("long.pdf", pages="1")
    assert first["pages_read"] == [1]
    assert len(first["text"]) < len(full["text"])


def test_bad_page_range_is_a_clear_error():
    tools.write_pdf("doc.pdf", title="T", body="body")
    with pytest.raises(tools.ToolError, match="page range"):
        tools.read_pdf("doc.pdf", pages="banana")


def test_empty_body_and_title_are_refused():
    with pytest.raises(tools.ToolError, match="title"):
        tools.write_pdf("doc.pdf", title="  ", body="text")
    with pytest.raises(tools.ToolError, match="empty"):
        tools.write_pdf("doc.pdf", title="T", body="   ")


def test_wrong_tool_for_the_file_type_says_which_tool_to_use():
    tools.write_excel("book.xlsx", rows=[["a"]])
    with pytest.raises(tools.ToolError, match="read_excel"):
        tools.read_pdf("book.xlsx")


# --- missing files and listing --------------------------------------------

def test_missing_file_error_lists_what_is_available():
    tools.write_excel("present.xlsx", rows=[["a"]])
    with pytest.raises(tools.ToolError, match="present.xlsx"):
        tools.read_excel("absent.xlsx")


def test_list_files_reports_what_was_written():
    tools.write_excel("book.xlsx", rows=[["a"]])
    tools.write_pdf("doc.pdf", title="T", body="body")
    names = {entry["name"] for entry in tools.list_files()["files"]}
    assert names == {"book.xlsx", "doc.pdf"}


# --- regressions from the second real session (1 Sep 2026) ----------------

def test_asking_for_the_xlsx_when_only_the_pdf_exists_names_the_right_tool():
    tools.write_pdf("phase04_invoice.pdf", title="Invoice", body="Total 22,640 AED")
    with pytest.raises(tools.ToolError) as caught:
        tools.read_excel("phase04_invoice.xlsx")
    message = str(caught.value)
    assert "phase04_invoice.pdf" in message
    assert "read_pdf" in message
    assert "Do not ask the user" in message


def test_reading_a_pdf_with_the_spreadsheet_tool_says_which_tool_to_use():
    tools.write_pdf("doc.pdf", title="T", body="body")
    with pytest.raises(tools.ToolError, match="read_pdf"):
        tools.read_excel("doc.pdf")
    # and the old misleading line about .xls is gone
    with pytest.raises(tools.ToolError) as caught:
        tools.read_excel("doc.pdf")
    assert ".xls format" not in str(caught.value)


def test_a_wrong_sheet_name_on_a_single_sheet_file_reads_it_anyway():
    tools.write_excel("book.xlsx", rows=[["a", 1]], sheet="Q1")
    result = tools.read_excel("book.xlsx", sheet="Sheet1")
    assert result["sheet"] == "Q1"
    assert result["rows"][0] == ["a", 1]
    assert "only one sheet" in result["note"]


def test_a_wrong_sheet_name_with_several_sheets_still_asks():
    tools.write_excel("book.xlsx", rows=[["a", 1]], sheet="Q1")
    tools.write_excel("book.xlsx", rows=[["b", 2]], sheet="Q2")
    with pytest.raises(tools.ToolError, match="Q1, Q2"):
        tools.read_excel("book.xlsx", sheet="Sheet1")


def test_a_correct_sheet_name_is_never_second_guessed():
    tools.write_excel("book.xlsx", rows=[["a", 1]], sheet="Q1")
    tools.write_excel("book.xlsx", rows=[["b", 2]], sheet="Q2")
    result = tools.read_excel("book.xlsx", sheet="Q2")
    assert result["sheet"] == "Q2"
    assert "note" not in result


def test_list_files_says_which_tool_reads_each_file():
    tools.write_pdf("invoice.pdf", title="T", body="body")
    tools.write_excel("sales.xlsx", rows=[["a", 1]])
    by_name = {f["name"]: f["read_with"] for f in tools.list_files()["files"]}
    assert by_name == {"invoice.pdf": "read_pdf", "sales.xlsx": "read_excel"}


# --------------------------------------------------------------------------
# The two worlds must not be confused for each other.
# In real use the model asked for "Report.xlsx" -- Report being a database
# table, not a file -- and the error listed eighteen unrelated filenames.
# It then spent the rest of the conversation searching the data folder and
# eventually denied being able to reach a database at all.
# --------------------------------------------------------------------------

def test_asking_for_a_table_as_a_file_points_at_the_database(monkeypatch, tmp_path):
    from app import db, tools

    monkeypatch.setattr(db, "is_configured", lambda: True)
    monkeypatch.setattr(
        db, "known_table_names", lambda *a, **k: {"report": 'public."Report"'}
    )

    with pytest.raises(tools.ToolError) as caught:
        tools.read_excel("Report.xlsx")

    message = str(caught.value)
    assert "not a file" in message
    assert "query_to_excel" in message
    assert 'public."Report"' in message
    # The old behaviour, and what sent it hunting: a wall of filenames.
    assert "phase03_sales.xlsx" not in message


def test_a_real_missing_file_still_gets_the_normal_error(monkeypatch):
    from app import db, tools

    monkeypatch.setattr(db, "known_table_names", lambda *a, **k: {"report": "x"})
    with pytest.raises(tools.ToolError) as caught:
        tools.read_excel("definitely_not_a_table.xlsx")
    assert "No file named" in str(caught.value)


def test_the_hint_never_breaks_the_error_path(monkeypatch):
    """A database that is down must not stop a missing file from being reported."""
    from app import db, tools

    def explode(*a, **k):
        raise RuntimeError("database is on fire")

    monkeypatch.setattr(db, "table_hint", explode)
    with pytest.raises(tools.ToolError) as caught:
        tools.read_excel("nope.xlsx")
    assert "No file named" in str(caught.value)


# --------------------------------------------------------------------------
# Regressions found on 3 Sep 2026 by reading the code back.
# --------------------------------------------------------------------------

def test_trailing_blank_rows_do_not_look_like_missing_data(tmp_path, monkeypatch):
    # openpyxl reports a row for any stray formatted cell at the bottom of a
    # sheet. Counting those made total_rows too high and truncated true for a
    # sheet that had in fact been read in full -- which the prompt tells the
    # model to treat as a signal that it is missing data.
    from openpyxl import Workbook

    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    book = Workbook()
    sheet = book.active
    sheet.append(["h1", "h2"])
    sheet.append([1, 2])
    sheet.append([3, 4])
    sheet.cell(row=9, column=2).value = None      # touched, empty
    book.save(tmp_path / "blanks.xlsx")

    result = tools.read_excel("blanks.xlsx")
    assert result["total_rows"] == 3
    assert result["rows_returned"] == 3
    assert result["truncated"] is False


def test_a_new_file_says_its_data_starts_on_row_one(tmp_path, monkeypatch):
    # ws.max_row is 1 for an empty sheet as well as for a sheet with one row,
    # so a guess made before appending was off by one on every new file.
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    made = tools.write_excel("fresh.xlsx", [["a", 1], ["b", 2]])
    assert made["started_at_row"] == 1
    appended = tools.write_excel("fresh.xlsx", [["c", 3]], mode="append")
    assert appended["started_at_row"] == 3


def test_a_file_the_assistant_writes_can_be_found_by_search(tmp_path, monkeypatch):
    # Uploads were indexed and files written by the tools were not, so the
    # assistant could create a spreadsheet and then fail to find it.
    from app import search

    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "INDEX_PATH", tmp_path / "idx.sqlite3")
    monkeypatch.setattr(config, "INDEX_DIR", tmp_path)
    monkeypatch.setattr(search, "INDEX_PATH", tmp_path / "idx.sqlite3")

    tools.write_excel("written.xlsx", [["Kryptonite", 7]], headers=["item", "qty"])
    assert search.search("Kryptonite")["count"] == 1
